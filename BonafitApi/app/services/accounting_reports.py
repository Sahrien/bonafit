from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors as colors_lib
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas import AccountingEntryOut, AccountingReportOut
from app.services.accounting import DISCLAIMER, DOCUMENT_DISCLAIMER

MONEY = '#,##0.00 €'


def build_xlsx(
    pyg: AccountingReportOut,
    income: AccountingReportOut,
    expense: AccountingReportOut,
    vat: AccountingReportOut,
) -> bytes:
    book = Workbook()
    _summary_sheet(book.active, pyg)
    _entries_sheet(book.create_sheet("Ingresos"), income.entries, pyg)
    _entries_sheet(book.create_sheet("Gastos"), expense.entries, pyg)
    _vat_sheet(book.create_sheet("IVA"), vat)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _header_rows(report: AccountingReportOut) -> list[list[object]]:
    period = f"{report.from_.date().isoformat()} — {report.to.date().isoformat()}"
    return [
        [report.legalName],
        [period],
        [DOCUMENT_DISCLAIMER],
        [DISCLAIMER],
        [],
    ]


def _summary_sheet(sheet, report: AccountingReportOut) -> None:
    sheet.title = "Resumen"
    rows = _header_rows(report)
    rows.append(["Concepto", "Importe"])
    kpis = report.kpis
    rows.extend(
        [
            ["Ingresos", kpis.income],
            ["Gastos", kpis.expense],
            ["Resultado", kpis.result],
            ["Ingresos cobrados", kpis.paidIncome],
            ["Por cobrar", kpis.pendingIncome],
            ["IVA repercutido", kpis.vatCollected],
            ["IVA soportado", kpis.vatDeductible],
            ["IVA neto", kpis.vatNet],
        ]
    )
    rows.append([])
    rows.append(["Categoría", "Tipo", "Importe"])
    for item in report.breakdown:
        kind = "Ingreso" if item.kind == "income" else "Gasto"
        rows.append([item.name, kind, item.amount])
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=14)
    for cell in sheet["B7":"B15"][0] if False else []:
        pass
    for row_index in range(7, 16):
        sheet.cell(row=row_index, column=2).number_format = MONEYP
    start = 18
    for offset in range(len(report.breakdown)):
        sheet.cell(row=start + offset, column=3).number_format = MONEYP
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 16


MONEYP = '#,##0.00'


def _entries_sheet(sheet, entries: list[AccountingEntryOut], report: AccountingReportOut) -> None:
    for row in _header_rows(report):
        sheet.append(row)
    sheet.append(["Fecha", "Concepto", "Categoría", "Contraparte", "Base", "IVA", "Total", "Estado", "Método"])
    for entry in entries:
        sheet.append(
            [
                entry.date.date().isoformat(),
                entry.concept,
                entry.categoryName,
                entry.counterpartyName,
                entry.netAmount,
                entry.vatAmount,
                entry.amount,
                entry.paymentStatus,
                entry.paymentMethod,
            ]
        )
    for row_index in range(7, 7 + len(entries)):
        for col in (5, 6, 7):
            sheet.cell(row=row_index, column=col).number_format = MONEYP
    sheet["A1"].font = Font(bold=True, size=14)
    for letter, width in {"A": 14, "B": 32, "C": 24, "D": 22, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12}.items():
        sheet.column_dimensions[letter].width = width


def _vat_sheet(sheet, report: AccountingReportOut) -> None:
    for row in _header_rows(report):
        sheet.append(row)
    sheet.append(["Base ingresos", "IVA repercutido", "Base gastos", "IVA soportado", "Cuota neta"])
    income_net = sum(item.netAmount for item in report.entries if item.type == "income")
    expense_net = sum(item.netAmount for item in report.entries if item.type == "expense")
    sheet.append(
        [
            round(income_net, 2),
            report.kpis.vatCollected,
            round(expense_net, 2),
            report.kpis.vatDeductible,
            report.kpis.vatNet,
        ]
    )
    for col in range(1, 6):
        sheet.cell(row=7, column=col).number_format = MONEYP
    sheet["A1"].font = Font(bold=True, size=14)


def build_pdf(
    report: AccountingReportOut,
    logo_path: Path | None = None,
    colors: tuple[str, str] | None = None,
) -> bytes:
    buffer = BytesIO()
    primary = _hex_color(colors[0] if colors else "#0f766e")
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("BonaTitle", parent=styles["Heading1"], textColor=primary, fontSize=16, spaceAfter=4)
    muted = ParagraphStyle("BonaMuted", parent=styles["Normal"], textColor=colors_lib.HexColor("#5c5854"), fontSize=9, spaceAfter=6)
    body = ParagraphStyle("BonaBody", parent=styles["Normal"], fontSize=10, spaceAfter=4)
    story: list = []
    if logo_path is not None:
        try:
            story.append(Image(str(logo_path), width=28 * mm, height=14 * mm, kind="proportional"))
            story.append(Spacer(1, 4 * mm))
        except Exception:
            pass
    story.append(Paragraph(report.legalName, title))
    story.append(Paragraph(report.title, styles["Heading2"]))
    period = f"{report.from_.date().isoformat()} — {report.to.date().isoformat()}"
    story.append(Paragraph(period, muted))
    story.append(Paragraph(DOCUMENT_DISCLAIMER, muted))
    story.append(Paragraph(DISCLAIMER, muted))
    kpis = report.kpis
    table_data = [
        ["Ingresos", _fmt(kpis.income), "Gastos", _fmt(kpis.expense)],
        ["Resultado", _fmt(kpis.result), "Por cobrar", _fmt(kpis.pendingIncome)],
        ["IVA repercutido", _fmt(kpis.vatCollected), "IVA soportado", _fmt(kpis.vatDeductible)],
    ]
    table = Table(table_data, colWidths=[40 * mm, 40 * mm, 40 * mm, 40 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), primary),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors_lib.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.3, colors_lib.HexColor("#d4d0cc")),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 6 * mm))
    if report.kind in {"income-book", "expense-book"} and report.entries:
        story.append(Paragraph("Movimientos", styles["Heading3"]))
        rows = [["Fecha", "Concepto", "Categoría", "Total"]]
        for entry in report.entries[:40]:
            rows.append(
                [
                    entry.date.date().isoformat(),
                    entry.concept[:40],
                    entry.categoryName[:24],
                    _fmt(entry.amount),
                ]
            )
        book = Table(rows, colWidths=[28 * mm, 62 * mm, 50 * mm, 28 * mm])
        book.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), primary),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors_lib.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors_lib.HexColor("#d4d0cc")),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(book)
    elif report.breakdown:
        story.append(Paragraph("Por categoría", styles["Heading3"]))
        rows = [["Categoría", "Tipo", "Importe"]]
        for item in report.breakdown:
            rows.append([item.name, "Ingreso" if item.kind == "income" else "Gasto", _fmt(item.amount)])
        chart = Table(rows, colWidths=[90 * mm, 35 * mm, 35 * mm])
        chart.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), primary),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors_lib.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors_lib.HexColor("#d4d0cc")),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(chart)
    else:
        story.append(Paragraph("No hay movimientos en este periodo.", body))
    doc.build(story)
    return buffer.getvalue()


def _fmt(value: float) -> str:
    return f"{value:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def _hex_color(value: str):
    try:
        return colors_lib.HexColor(value)
    except Exception:
        return colors_lib.HexColor("#0f766e")
